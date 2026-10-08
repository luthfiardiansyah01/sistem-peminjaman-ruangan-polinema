"""
Generator dokumen bukti revisi before/after.
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

BASE_FONT = "Times New Roman"

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = BASE_FONT
    r.font.size = Pt(16)
    r.bold = True
    p.paragraph_format.space_after = Pt(12)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = BASE_FONT
    r.font.size = Pt(13)
    r.bold = True
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after  = Pt(6)
    return p

def add_h3(doc, text, color_hex="1F497D"):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = BASE_FONT
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color_hex)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(4)
    return p

def add_body(doc, text, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = BASE_FONT
    r.font.size = Pt(11)
    r.italic = italic
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = Pt(16)
    return p

def add_colored_box(doc, text, fill_hex, text_hex="000000"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.rows[0].cells[0]
    cell.text = text
    # fill
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill_hex)
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)
    for para in cell.paragraphs:
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            run.font.name = BASE_FONT
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor.from_string(text_hex)
    doc.add_paragraph()

def add_comparison(doc, before_text, after_text):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = 'Table Grid'

    # Header
    hdr_row = tbl.rows[0]
    hdr_row.cells[0].text = "SEBELUM (Before)"
    hdr_row.cells[1].text = "SESUDAH (After)"
    for i, cell in enumerate(hdr_row.cells):
        fill = "FF6B6B" if i == 0 else "51CF66"
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), fill)
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:val'), 'clear')
        tcPr.append(shd)
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.name = BASE_FONT
                run.font.size = Pt(11)
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("FFFFFF")

    # Content row
    data_row = tbl.add_row()
    data_row.cells[0].text = before_text
    data_row.cells[1].text = after_text
    for i, cell in enumerate(data_row.cells):
        fill = "FFF5F5" if i == 0 else "F0FFF4"
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), fill)
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:val'), 'clear')
        tcPr.append(shd)
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.name = BASE_FONT
                run.font.size = Pt(10)
    doc.add_paragraph()

# ─────────────────────────────────────────────────────────────────────────────
doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

# Cover
add_h1(doc, "DOKUMENTASI BUKTI REVISI")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Sistem Peminjaman Ruangan (SPR) JTI Polinema")
r.font.name = BASE_FONT
r.font.size = Pt(13)
r.bold = True

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Laporan Pemenuhan Revisi Penulisan Laporan Skripsi")
r2.font.name = BASE_FONT
r2.font.size = Pt(12)

doc.add_paragraph()
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = p3.add_run("Tanggal Revisi: 6 Oktober 2026")
r3.font.name = BASE_FONT
r3.font.size = Pt(11)
r3.italic = True

doc.add_paragraph()

# Summary table
add_h2(doc, "Ringkasan Revisi")
tbl_sum = doc.add_table(rows=1, cols=3)
tbl_sum.style = 'Table Grid'
hdr = tbl_sum.rows[0]
hdr.cells[0].text = "No"
hdr.cells[1].text = "Poin Revisi"
hdr.cells[2].text = "Status"
for cell in hdr.cells:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), '1F497D')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.name = BASE_FONT
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("FFFFFF")

summary_items = [
    ("1", "Penulisan disesuaikan dengan template bahan_skripsi (BAB I–VII, Times New Roman, margin 4-3-3-3)", "SELESAI"),
    ("2", "Gaya bahasa disesuaikan dengan draft proposal PDF (formal akademik Bahasa Indonesia)", "SELESAI"),
    ("3", "Cover, Pengesahan, Pernyataan, Kata Pengantar tidak diubah (klien yang handle)", "TIDAK DIUBAH (sesuai instruksi)"),
    ("4", "Diagram disesuaikan gaya klien: Use Case, Use Case Utama, Activity Diagram, ERD (.drawio)", "SELESAI"),
    ("5", "Mockup menggunakan aset dari File-Figma-SPR-JTI (direferensikan dalam dokumen)", "SELESAI"),
]
for no, item, status in summary_items:
    row = tbl_sum.add_row()
    row.cells[0].text = no
    row.cells[1].text = item
    row.cells[2].text = status
    fill_status = "E8F5E9" if "SELESAI" in status else "FFF9C4"
    tcPr = row.cells[2]._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill_status)
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)
    for i, cell in enumerate(row.cells):
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.name = BASE_FONT
                run.font.size = Pt(10)

doc.add_paragraph()

# ─── POIN 1 & 2: Dokumen Skripsi ─────────────────────────────────────────────
doc.add_page_break()
add_h2(doc, "POIN 1 & 2: Penulisan & Gaya Bahasa Laporan")
add_h3(doc, "Perubahan yang Dilakukan")
add_body(doc, (
    "Dokumen laporan sebelumnya berupa 'Reverse Engineering Report' dalam format teknis singkat "
    "(2 halaman) yang tidak sesuai format skripsi. Revisi mengubah dokumen menjadi laporan skripsi "
    "formal dengan struktur BAB I–VII mengikuti template bahan_skripsi, dengan konten yang "
    "disesuaikan dari draft proposal PDF referensi klien."
))

add_comparison(doc,
    "SEBELUM:\n"
    "File: Reverse_Engineering_Report_SPR_JTI.docx\n"
    "- Format: Technical reverse engineering report\n"
    "- Jumlah halaman: ~2 halaman\n"
    "- Struktur: Summary, Components, Flow, Data, Integration, Security, References\n"
    "- Gaya: Teknis Inggris, tanpa format akademik\n"
    "- Font: Default (Calibri)\n"
    "- Tidak ada BAB, sub-bab, atau format skripsi\n"
    "- Tidak ada Abstrak, Daftar Pustaka\n"
    "- Tidak ada kebutuhan fungsional terstruktur",

    "SESUDAH:\n"
    "File: revisi_output/Skripsi_SPR_JTI_Revised.docx\n"
    "- Format: Skripsi Diploma IV Polinema\n"
    "- Struktur: Abstrak, BAB I (Pendahuluan), BAB II (Landasan Teori), BAB III (Metode), BAB IV (Analisis & Perancangan), BAB V–VII (placeholder)\n"
    "- Font: Times New Roman 12pt\n"
    "- Margin: 4-3-3-3 cm\n"
    "- Spasi: 1.5 (18pt line spacing)\n"
    "- Gaya bahasa: formal akademik Bahasa Indonesia\n"
    "- Tabel 20 kebutuhan fungsional (F-01 s.d. F-20)\n"
    "- Tabel ERD 20 tabel basis data\n"
    "- Daftar Pustaka 11 referensi akademik"
)

add_h3(doc, "Konten BAB yang Dibuat")
items_bab = [
    "BAB I: Latar Belakang, Rumusan Masalah, Batasan Masalah, Tujuan, Manfaat",
    "BAB II: Sistem Informasi, Framework Laravel, Metodologi Waterfall, UML, ERD, Black Box Testing, UEQ, Penelitian Terkait",
    "BAB III: Jenis Penelitian, Metodologi Waterfall (4 fase), 20 Kebutuhan Fungsional (tabel), Kebutuhan Non-Fungsional, Alat dan Bahan",
    "BAB IV: Analisis Sistem Berjalan, Sistem Diusulkan, Use Case Diagram, Use Case Utama, Activity Diagram, ERD, Tabel ERD (20 tabel), Perancangan Mockup",
    "BAB V–VII: Placeholder (menunggu hasil pengujian aktual)",
    "Abstrak & Daftar Pustaka: Terisi penuh",
]
for item in items_bab:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(item)
    r.font.name = BASE_FONT
    r.font.size = Pt(11)

# ─── POIN 4: Diagram ─────────────────────────────────────────────────────────
doc.add_page_break()
add_h2(doc, "POIN 4: Diagram (Use Case, Activity Diagram, ERD)")
add_h3(doc, "Perubahan yang Dilakukan")
add_body(doc, (
    "Sebelumnya tidak ada diagram dalam laporan. Revisi menghasilkan 4 file .drawio baru yang "
    "mengikuti gaya diagram klien dari file referensi 'UML Sistem Peminjaman Ruangan JTI.drawio', "
    "yaitu: strokeWidth=2, orthogonalEdgeStyle, swimlane, dan flowchart shapes dari mxgraph."
))

add_comparison(doc,
    "SEBELUM:\n"
    "- Tidak ada diagram sama sekali dalam laporan\n"
    "- Tidak ada file .drawio\n"
    "- Tidak ada visualisasi use case, activity flow, atau ERD",

    "SESUDAH:\n"
    "File di revisi_output/diagrams/:\n"
    "1. Use_Case_Diagram.drawio\n"
    "   - 4 aktor: Admin, Peminjam, Verifikator, Sistem\n"
    "   - 20 use case sesuai kebutuhan fungsional\n"
    "   - System boundary swimlane\n\n"
    "2. Use_Case_Fungsionalitas_Utama.drawio\n"
    "   - Fokus alur peminjaman inti\n"
    "   - <<include>> dan <<extend>> relationships\n"
    "   - 4 grup: Pengajuan, Verifikasi, Penyelesaian, Verifikasi Penyelesaian\n\n"
    "3. Activity_Diagram.drawio\n"
    "   - 4 swimlane: Peminjam, Verifikator, Admin, Sistem\n"
    "   - Alur lengkap dari login hingga selesai\n"
    "   - Flowchart shapes: start/end, proses, keputusan, I/O\n\n"
    "4. ERD.drawio\n"
    "   - 20 tabel basis data\n"
    "   - Relasi 1:1 dan 1:N dengan ER arrow style\n"
    "   - Color-coded per kategori tabel"
)

add_h3(doc, "Gaya Diagram Klien yang Diterapkan")
styles = [
    ("Swimlane header",   "style=\"swimlane;whiteSpace=wrap;html=1;\""),
    ("Proses (rounded)",  "style=\"rounded=1;whiteSpace=wrap;html=1;absoluteArcSize=1;arcSize=14;strokeWidth=2;\""),
    ("Keputusan (diamond)","style=\"strokeWidth=2;html=1;shape=mxgraph.flowchart.decision;whiteSpace=wrap;\""),
    ("I/O (parallelogram)","style=\"shape=parallelogram;html=1;strokeWidth=2;perimeter=parallelogramPerimeter;whiteSpace=wrap;rounded=1;arcSize=12;size=0.23;\""),
    ("Start/End",          "style=\"strokeWidth=2;html=1;shape=mxgraph.flowchart.start_1;whiteSpace=wrap;\""),
    ("Edge (alur)",        "style=\"edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;strokeWidth=2;\""),
]
tbl_style = doc.add_table(rows=1, cols=2)
tbl_style.style = 'Table Grid'
hdr_st = tbl_style.rows[0]
hdr_st.cells[0].text = "Elemen"
hdr_st.cells[1].text = "Style yang Diterapkan"
for cell in hdr_st.cells:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), '3E7CBB')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.name = BASE_FONT
            run.font.size = Pt(10)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
for elem, style in styles:
    row = tbl_style.add_row()
    row.cells[0].text = elem
    row.cells[1].text = style
    for cell in row.cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9)

doc.add_paragraph()

# ─── POIN 5: Mockup ──────────────────────────────────────────────────────────
doc.add_page_break()
add_h2(doc, "POIN 5: Mockup Antarmuka")
add_h3(doc, "Perubahan yang Dilakukan")
add_body(doc, (
    "Sebelumnya tidak ada mockup dalam laporan. Revisi mengintegrasikan referensi mockup dari "
    "aset Figma klien (folder File-Figma-SPR-JTI, 364 file PNG) ke dalam dokumen skripsi di "
    "BAB IV bagian Perancangan Antarmuka (Seksi 4.7). Mockup direferensikan dengan caption gambar "
    "sesuai halaman-halaman utama sistem."
))

add_comparison(doc,
    "SEBELUM:\n"
    "- Tidak ada mockup dalam laporan\n"
    "- Tidak ada referensi antarmuka pengguna\n"
    "- Tidak ada perancangan UI/UX terdokumentasi",

    "SESUDAH (BAB IV Seksi 4.7):\n"
    "Mockup direferensikan untuk:\n"
    "- Halaman Login (Login Admin.png)\n"
    "- Halaman Dasbor Admin (Dashboard Admin.png)\n"
    "- Halaman Form Pengajuan Peminjaman\n"
    "- Daftar Jadwal (Daftar Jadwal.png)\n"
    "- Daftar Ruangan (Daftar Ruangan.png)\n"
    "- Tambah/Edit/Hapus/Lihat Data Jadwal\n"
    "- Tambah/Edit/Hapus/Lihat Data Ruangan\n\n"
    "Aset PNG tersedia di:\n"
    "bahan_skripsi/File-Figma-SPR-JTI/ (364 file PNG)"
)

add_h3(doc, "Daftar File Mockup yang Tersedia")
mockup_files = [
    "Login Admin.png", "Dashboard Admin.png", "Daftar Jadwal.png",
    "Daftar Ruangan.png", "Tambah Data Jadwal.png", "Lihat Data Jadwal.png",
    "Edit Data Jadwal.png", "Hapus Data Jadwal.png", "Berhasil Tambah Jadwal.png",
    "Berhasil Edit Jadwal.png", "Berhasil Hapus Jadwal.png",
    "Tambah Data Ruangan.png", "Edit Data Ruangan.png", "Hapus Data Ruangan.png",
    "Lihat Data Ruangan.png", "Daftar Pengguna Admin.png", "ADMIN.png",
    "... (total 364 file PNG dari Figma export)",
]
for f in mockup_files:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(f)
    r.font.name = "Courier New"
    r.font.size = Pt(10)

# ─── POIN 3: Catatan ─────────────────────────────────────────────────────────
doc.add_paragraph()
add_h2(doc, "POIN 3: Halaman Cover, Pengesahan, Pernyataan, Kata Pengantar")
add_colored_box(doc,
    "STATUS: TIDAK DIUBAH — Sesuai instruksi klien.\n"
    "Klien menyatakan bahwa bagian ini akan dimodifikasi sendiri oleh klien. "
    "Tidak ada perubahan yang dilakukan pada halaman Cover, Pengesahan, Pernyataan, dan Kata Pengantar.",
    "FFF9C4"
)

# ─── Daftar File Output ───────────────────────────────────────────────────────
doc.add_page_break()
add_h2(doc, "Daftar File Output Revisi")
add_body(doc, "Seluruh file hasil revisi tersimpan di folder revisi_output/ dalam direktori proyek:")

output_files = [
    ("revisi_output/Skripsi_SPR_JTI_Revised.docx",
     "Dokumen skripsi revisi utama (BAB I–VII, Abstrak, Daftar Pustaka)"),
    ("revisi_output/diagrams/Use_Case_Diagram.drawio",
     "Use Case Diagram – 4 aktor, 20 use case, system boundary"),
    ("revisi_output/diagrams/Use_Case_Fungsionalitas_Utama.drawio",
     "Use Case Utama – alur peminjaman inti dengan <<include>>/<<extend>>"),
    ("revisi_output/diagrams/Activity_Diagram.drawio",
     "Activity Diagram – 4 swimlane, alur lengkap peminjaman"),
    ("revisi_output/diagrams/ERD.drawio",
     "ERD – 20 tabel basis data dengan relasi 1:1 dan 1:N"),
    ("revisi_output/Bukti_Revisi.docx",
     "Dokumen ini – bukti before/after seluruh poin revisi"),
]

tbl_out = doc.add_table(rows=1, cols=2)
tbl_out.style = 'Table Grid'
hdr_out = tbl_out.rows[0]
hdr_out.cells[0].text = "File"
hdr_out.cells[1].text = "Deskripsi"
for cell in hdr_out.cells:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), '2E7D32')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.name = BASE_FONT
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("FFFFFF")

for fpath, fdesc in output_files:
    row = tbl_out.add_row()
    row.cells[0].text = fpath
    row.cells[1].text = fdesc
    for i, cell in enumerate(row.cells):
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.name = "Courier New" if i == 0 else BASE_FONT
                run.font.size = Pt(10)

doc.add_paragraph()
add_body(doc,
    "Catatan: Mockup (aset PNG) sudah tersedia di bahan_skripsi/File-Figma-SPR-JTI/ dan "
    "direferensikan dalam dokumen skripsi pada BAB IV. File PNG tidak disalin ke revisi_output "
    "karena ukurannya besar (364 file).",
    italic=True
)

# ─── Save ─────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(__file__), "Bukti_Revisi.docx")
doc.save(out_path)
print(f"[OK] Bukti revisi tersimpan: {out_path}")
