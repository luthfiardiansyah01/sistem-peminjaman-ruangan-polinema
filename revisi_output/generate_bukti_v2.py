"""
Generator dokumen Bukti_Revisi_v2.docx
Perbandingan before/after berbasis template asli SPR JTI.
"""
import sys, os
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = r'revisi_output\Bukti_Revisi_v2.docx'
doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

F = 'Times New Roman'

# ── Helpers ──────────────────────────────────────────────────────────────────

def shd(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:fill'), hex_color)
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:val'), 'clear')
    tcPr.append(s)

def p(text, size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None, space_after=4):
    para = doc.add_paragraph()
    para.alignment = align
    run = para.add_run(text)
    run.font.name  = F
    run.font.size  = Pt(size)
    run.bold       = bold
    run.italic     = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    para.paragraph_format.space_after  = Pt(space_after)
    para.paragraph_format.line_spacing = Pt(15)
    return para

def h1(text):
    p(text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

def h2(text, color='1F497D'):
    pa = doc.add_paragraph()
    run = pa.add_run(text)
    run.font.name = F; run.font.size = Pt(13); run.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    pa.paragraph_format.space_before = Pt(10)
    pa.paragraph_format.space_after  = Pt(4)

def h3(text):
    pa = doc.add_paragraph()
    run = pa.add_run(text)
    run.font.name = F; run.font.size = Pt(11); run.bold = True
    pa.paragraph_format.space_before = Pt(6)
    pa.paragraph_format.space_after  = Pt(3)

def separator():
    pa = doc.add_paragraph()
    pa.paragraph_format.space_after = Pt(6)
    run = pa.add_run('─' * 90)
    run.font.name = F; run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string('AAAAAA')

def badge(text, fill='2E7D32', text_color='FFFFFF'):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    shd(cell, fill)
    para = cell.paragraphs[0]
    run = para.add_run(f'  {text}  ')
    run.font.name = F; run.font.size = Pt(10); run.bold = True
    run.font.color.rgb = RGBColor.from_string(text_color)
    para.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def before_after(before_text, after_text):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.style = 'Table Grid'
    # Headers
    for i, (label, fill) in enumerate([('SEBELUM (Template Asli)', 'C62828'), ('SESUDAH (Konten Aktual)', '1B5E20')]):
        cell = tbl.rows[0].cells[i]
        shd(cell, fill)
        para = cell.paragraphs[0]
        run = para.add_run(label)
        run.font.name = F; run.font.size = Pt(11); run.bold = True
        run.font.color.rgb = RGBColor.from_string('FFFFFF')
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Content
    fills = ['FFF8F8', 'F1F8E9']
    for i, (text, fill_hex) in enumerate([(before_text, fills[0]), (after_text, fills[1])]):
        cell = tbl.rows[1].cells[i]
        shd(cell, fill_hex)
        cell.text = text
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in para.runs:
                run.font.name = F
                run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def stats_table(rows_data, headers):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = 'Table Grid'
    hdr = tbl.rows[0]
    for i, h in enumerate(headers):
        shd(hdr.cells[i], '1565C0')
        para = hdr.cells[i].paragraphs[0]
        run = para.add_run(h)
        run.font.name = F; run.font.size = Pt(10); run.bold = True
        run.font.color.rgb = RGBColor.from_string('FFFFFF')
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_data in rows_data:
        row = tbl.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = val
            for para in row.cells[i].paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.name = F; run.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ════════════════════════════════════════════════════════════════════════════
# HALAMAN COVER
# ════════════════════════════════════════════════════════════════════════════
h1('DOKUMENTASI BUKTI REVISI')
p('Sistem Peminjaman Ruangan (SPR) JTI Polinema', size=13, bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
p('Pemenuhan Revisi Penulisan Laporan Skripsi', size=11,
  align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, space_after=4)
p('Tanggal Revisi: 6 Oktober 2026', size=10, italic=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

separator()

# ── Ringkasan ──────────────────────────────────────────────────────────────
h2('RINGKASAN REVISI')
stats_table([
    ('1', 'Penulisan sesuai template asli bahan_skripsi (BAB I–VII, style Heading1/2, Normal)', 'SELESAI ✓'),
    ('2', 'Gaya bahasa formal akademik Bahasa Indonesia sesuai draft PDF referensi klien', 'SELESAI ✓'),
    ('3', 'Cover, Pengesahan, Pernyataan, Kata Pengantar TIDAK diubah', 'TIDAK DIUBAH (sesuai instruksi)'),
    ('4', 'Diagram: Use Case, Use Case Utama, Activity Diagram, ERD (.drawio, gaya klien)', 'SELESAI ✓'),
    ('5', 'Mockup direferensikan dari aset Figma klien di BAB IV Seksi 4.7', 'SELESAI ✓'),
], ['No', 'Poin Revisi', 'Status'])

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# POIN 1 & 2 — DOKUMEN SKRIPSI
# ════════════════════════════════════════════════════════════════════════════
h2('POIN 1 & 2 — PENULISAN & GAYA BAHASA LAPORAN')

h3('File yang Direvisi')
p('• Template asli: bahan_skripsi/Skripsi Pengembangan Sistem Peminjaman Ruangan JTI Polinema berbasis Web.docx')
p('• Output baru:   revisi_output/Skripsi_SPR_JTI_v2.docx')
doc.add_paragraph()

h3('Perubahan: Halaman Abstrak & Abstract')
before_after(
    'SEBELUM (isi contoh dari skripsi lain):\n\n'
    'ABSTRAK\n'
    'Kartika P., Anggi. "Pengembangan Aplikasi Manajemen\n'
    'Stok UMKM dengan Fitur Prediksi Penjualan..."\n\n'
    'Pada proses bisnis UMKM di Indonesia...\n'
    'Lorem ipsum dolor sit amet. A quick brown fox\n'
    'jumps over a lazy frog. [PLACEHOLDER]\n\n'
    'Kata Kunci: Sistem Informasi, Jaringan Syaraf\n'
    'Tiruan, UMKM',

    'SESUDAH (konten aktual SPR JTI):\n\n'
    'ABSTRAK\n'
    'Paramita, Ratih. "Pengembangan Sistem Peminjaman\n'
    'Ruangan JTI Polinema Berbasis Web."\n\n'
    'Pengelolaan peminjaman ruangan di JTI Polinema\n'
    'yang masih manual menimbulkan bentrok jadwal,\n'
    'proses verifikasi lambat...\n\n'
    'Metodologi: Waterfall, Laravel 10, PHP 8.1, MySQL.\n\n'
    'Hasil: 20 fungsionalitas, Black Box ✓, UEQ positif.\n\n'
    'Kata Kunci: Sistem Peminjaman Ruangan, Laravel 10,\n'
    'Waterfall, Verifikasi Multi-Tingkat, Black Box, UEQ'
)

h3('Perubahan: BAB I — Pendahuluan')
before_after(
    'SEBELUM:\n\n'
    'BAB I. PENDAHULUAN  [List Paragraph]\n\n'
    'Sub Bab  [placeholder]\n\n'
    '"Beberapa uraian yang ada pada bagian\n'
    'pendahuluan adalah latar belakang, rumusan\n'
    'masalah, tujuan, batasan masalah..."\n\n'
    '"Ini adalah teks sub bab 1.1 di mana margin\n'
    'kiri dimulai persis di bawah teks sub bab,\n'
    'namun pada alinea pertama menjorok..."\n\n'
    '[isi contoh lorem ipsum & panduan format]',

    'SESUDAH:\n\n'
    'BAB I  [Heading 1]\n'
    'PENDAHULUAN  [Heading 1]\n\n'
    '1.1 Latar Belakang  [Heading 2]\n'
    'Politeknik Negeri Malang (Polinema)...\n'
    '[4 paragraf konten aktual]\n\n'
    '1.2 Rumusan Masalah  [Heading 2]\n'
    '1.3 Batasan Masalah  [Heading 2]\n'
    '1.4 Tujuan Penelitian  [Heading 2]\n'
    '1.5 Manfaat Penelitian  [Heading 2]\n'
    '  1.5.1 Manfaat Teoritis  [Heading 2]\n'
    '  1.5.2 Manfaat Praktis  [Heading 2]'
)

h3('Perubahan: BAB II — Landasan Teori')
before_after(
    'SEBELUM:\n\n'
    'BAB II. LANDASAN TEORI  [List Paragraph]\n\n'
    'Sub Bab  [placeholder]\n\n'
    '"Landasan Teori berisikan teori-teori yang\n'
    'relevan yang melengkapi latar belakang dan\n'
    'dijadikan referensi dalam penelitian..."\n\n'
    '"Berikut ini adalah contoh kutipan pernyataan\n'
    'yang berasal dari 2 penulis..."\n\n'
    '[contoh format kutipan saja, tidak ada konten]',

    'SESUDAH:\n\n'
    'BAB II  [Heading 1]\n'
    'LANDASAN TEORI  [Heading 1]\n\n'
    '2.1 Sistem Informasi\n'
    '2.2 Peminjaman Ruangan\n'
    '2.3 Framework Laravel\n'
    '2.4 Metodologi Waterfall\n'
    '2.5 Unified Modeling Language (UML)\n'
    '2.6 Entity Relationship Diagram (ERD)\n'
    '2.7 Black Box Testing\n'
    '2.8 User Experience Questionnaire (UEQ)\n'
    '2.9 Penelitian Terkait\n\n'
    '[9 subbab dengan konten akademik aktual,\n'
    'masing-masing 1–2 paragraf dengan sitasi]'
)

h3('Perubahan: BAB III — Metodologi Pengembangan')
before_after(
    'SEBELUM:\n\n'
    'BAB III. METODOLOGI PENGEMBANGAN\n\n'
    'Sub Bab  [placeholder]\n\n'
    '"Pada bab ini Terdiri dari langkah-langkah\n'
    'yang akan membimbing penulis memilih metode,\n'
    'teknik, prosedur yang dapat digunakan..."\n\n'
    '[hanya panduan isi, tidak ada tabel\n'
    'kebutuhan fungsional atau diagram]',

    'SESUDAH:\n\n'
    'BAB III  [Heading 1]\n'
    'METODOLOGI PENGEMBANGAN  [Heading 1]\n\n'
    '3.1 Jenis Penelitian\n'
    '3.2 Metode Pengembangan Sistem (Waterfall)\n'
    '  3.2.1 Analisis Kebutuhan\n'
    '  3.2.2 Perancangan Sistem\n'
    '  3.2.3 Implementasi\n'
    '  3.2.4 Pengujian\n'
    '3.3 Kebutuhan Fungsional Sistem\n'
    '  → Tabel 20 kebutuhan F-01 s.d. F-20\n'
    '3.4 Kebutuhan Non-Fungsional Sistem\n'
    '3.5 Alat dan Bahan Penelitian'
)

h3('Perubahan: BAB IV — Analisis & Perancangan')
before_after(
    'SEBELUM:\n\n'
    'BAB IV. ANALISIS DAN PERANCANGAN SISTEM\n\n'
    'Sub Bab  [placeholder]\n\n'
    '"Pada bab ini terangkanlah proses-proses\n'
    'sebelum membuat sistem yang meliputi,\n'
    'namun tidak terbatas pada: analisa\n'
    'kebutuhan, use case diagram..."\n\n'
    '"Pada bagian ini diuraikan dengan jelas\n'
    'sistem yang akan dibuat dan kebutuhan\n'
    'sistem..."\n\n'
    '[hanya panduan isi, tidak ada diagram,\n'
    'ERD, atau mockup]',

    'SESUDAH:\n\n'
    'BAB IV  [Heading 1]\n'
    'ANALISIS DAN PERANCANGAN SISTEM  [Heading 1]\n\n'
    '4.1 Analisis Sistem Berjalan\n'
    '4.2 Analisis Sistem yang Diusulkan\n'
    '4.3 Use Case Diagram → [Gambar 4.1]\n'
    '4.4 Use Case Fungsionalitas Utama → [Gambar 4.2]\n'
    '4.5 Activity Diagram → [Gambar 4.3]\n'
    '4.6 Perancangan Basis Data\n'
    '  4.6.1 ERD → [Gambar 4.4]\n'
    '  4.6.2 Tabel ERD → Tabel 20 tabel DB\n'
    '4.7 Perancangan Antarmuka (Mockup)\n'
    '  4.7.1–4.7.5 Login, Dasbor, Pengajuan,\n'
    '              Jadwal, Ruangan → [Gambar 4.5–4.9]'
)

h3('Perubahan: BAB V–VII & Daftar Pustaka')
before_after(
    'SEBELUM:\n\n'
    'BAB V. IMPLEMENTASI DAN PENGUJIAN\n'
    'BAB VI. HASIL DAN PEMBAHASAN\n'
    'BAB VII. KESIMPULAN DAN SARAN\n'
    'DAFTAR PUSTAKA\n\n'
    '[semua berisi placeholder panduan penulisan\n'
    'dan contoh format dari skripsi lain:\n'
    '"Tuliskan kesimpulan dari penelitian Anda...\n'
    '"Baxter, C. (1997). Race equality in health\n'
    'care and education..." — tidak relevan]',

    'SESUDAH:\n\n'
    'BAB V: Implementasi Sistem (BD + antarmuka),\n'
    '  Tabel Black Box 10 skenario uji,\n'
    '  Pengujian UEQ\n\n'
    'BAB VI: Pembahasan hasil pengembangan,\n'
    '  Black Box, UEQ, fungsionalitas utama\n\n'
    'BAB VII: 4 kesimpulan + 4 saran\n'
    '  (relevan dengan tujuan penelitian)\n\n'
    'Daftar Pustaka: 11 referensi akademik\n'
    '  (Booch, Pressman, Sommerville, Laravel\n'
    '   Docs, Laugwitz UEQ, Myers, dll.)'
)

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# POIN 3 — TIDAK DIUBAH
# ════════════════════════════════════════════════════════════════════════════
h2('POIN 3 — HALAMAN YANG TIDAK DIUBAH')
badge('STATUS: TIDAK DIUBAH — Sesuai instruksi klien', fill='E65100')
p('Bagian-bagian berikut dipertahankan PERSIS seperti template asli, tidak ada modifikasi:')
for item in [
    'Halaman Cover (judul, nama, NIM, prodi, jurusan, tahun)',
    'Halaman Pengesahan',
    'Halaman Pernyataan',
    'Kata Pengantar',
    'Daftar Isi, Daftar Gambar, Daftar Tabel',
]:
    p(f'  ✓  {item}', size=11)

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# POIN 4 — DIAGRAM
# ════════════════════════════════════════════════════════════════════════════
h2('POIN 4 — DIAGRAM (.drawio)')

h3('Perbandingan sebelum dan sesudah')
before_after(
    'SEBELUM:\n\n'
    '• Tidak ada file diagram sama sekali\n'
    '• Tidak ada Use Case Diagram\n'
    '• Tidak ada Activity Diagram\n'
    '• Tidak ada ERD\n'
    '• BAB IV hanya berisi panduan teks\n'
    '  tanpa referensi gambar apapun',

    'SESUDAH (4 file .drawio):\n\n'
    '1. Use_Case_Diagram.drawio\n'
    '   4 aktor, 20 use case, system boundary\n\n'
    '2. Use_Case_Fungsionalitas_Utama.drawio\n'
    '   Alur peminjaman inti + <<include>>/<<extend>>\n\n'
    '3. Activity_Diagram.drawio\n'
    '   4 swimlane: Peminjam/Verifikator/Admin/Sistem\n\n'
    '4. ERD.drawio\n'
    '   20 tabel basis data dengan relasi 1:1 & 1:N'
)

h3('Gaya Diagram yang Diterapkan (sesuai file referensi klien)')
stats_table([
    ('Swimlane',    'swimlane;whiteSpace=wrap;html=1;strokeWidth=2;',                                                         'Pembatas aktor/area'),
    ('Proses',      'rounded=1;whiteSpace=wrap;html=1;arcSize=14;strokeWidth=2;',                                            'Aktivitas/proses'),
    ('Keputusan',   'strokeWidth=2;html=1;shape=mxgraph.flowchart.decision;whiteSpace=wrap;',                                'Percabangan Yes/No'),
    ('I/O',         'shape=parallelogram;html=1;strokeWidth=2;perimeter=parallelogramPerimeter;arcSize=12;size=0.23;',       'Input/Output'),
    ('Start/End',   'strokeWidth=2;html=1;shape=mxgraph.flowchart.start_1;whiteSpace=wrap;',                                 'Titik awal/akhir'),
    ('Alur/Edge',   'edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;strokeWidth=2;',        'Panah koneksi'),
], ['Elemen', 'Style (draw.io XML)', 'Fungsi'])

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# POIN 5 — MOCKUP
# ════════════════════════════════════════════════════════════════════════════
h2('POIN 5 — MOCKUP ANTARMUKA')

h3('Perbandingan sebelum dan sesudah')
before_after(
    'SEBELUM:\n\n'
    '• Tidak ada mockup dalam laporan\n'
    '• Tidak ada seksi perancangan antarmuka\n'
    '• BAB IV tidak menyebut UI/UX apapun\n'
    '• File Figma belum direferensikan',

    'SESUDAH (BAB IV Seksi 4.7):\n\n'
    '4.7 Perancangan Antarmuka (Mockup)\n'
    '  4.7.1 Halaman Login → [Gambar 4.5]\n'
    '  4.7.2 Halaman Dasbor Admin → [Gambar 4.6]\n'
    '  4.7.3 Halaman Pengajuan → [Gambar 4.7]\n'
    '  4.7.4 Halaman Daftar Jadwal → [Gambar 4.8]\n'
    '  4.7.5 Halaman Daftar Ruangan → [Gambar 4.9]\n\n'
    'Aset PNG tersedia di:\n'
    'bahan_skripsi/File-Figma-SPR-JTI/ (364 file)'
)

h3('Daftar Aset Mockup Tersedia (sampel dari 364 file PNG)')
stats_table([
    ('Login Admin.png',                'Halaman autentikasi pengguna'),
    ('Dashboard Admin.png',            'Dasbor statistik Admin'),
    ('Daftar Jadwal.png',              'Daftar jadwal penggunaan ruangan'),
    ('Daftar Ruangan.png',             'Daftar ruangan & ketersediaan'),
    ('Tambah Data Jadwal.png',         'Form tambah jadwal'),
    ('Lihat Data Jadwal.png',          'Detail data jadwal'),
    ('Edit Data Jadwal.png',           'Form edit jadwal'),
    ('Hapus Data Jadwal.png',          'Konfirmasi hapus jadwal'),
    ('Tambah Data Ruangan.png',        'Form tambah ruangan'),
    ('Edit Data Ruangan.png',          'Form edit ruangan'),
    ('Daftar Pengguna Admin.png',      'Daftar manajemen pengguna'),
    ('... dan 353 file lainnya',       '(total 364 file PNG dari Figma export)'),
], ['Nama File', 'Keterangan'])

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# DAFTAR FILE OUTPUT
# ════════════════════════════════════════════════════════════════════════════
h2('DAFTAR LENGKAP FILE OUTPUT REVISI')
p('Semua file tersimpan di folder revisi_output/ dalam direktori proyek:', space_after=6)

stats_table([
    ('1', 'revisi_output/Skripsi_SPR_JTI_v2.docx',
     'Skripsi revisi FINAL — berbasis template asli, konten aktual BAB I–VII (+1.4 MB)'),
    ('2', 'revisi_output/diagrams/Use_Case_Diagram.drawio',
     'Use Case Diagram — 4 aktor, 20 use case, system boundary'),
    ('3', 'revisi_output/diagrams/Use_Case_Fungsionalitas_Utama.drawio',
     'Use Case Utama — alur peminjaman inti dengan <<include>>/<<extend>>'),
    ('4', 'revisi_output/diagrams/Activity_Diagram.drawio',
     'Activity Diagram — 4 swimlane, alur lengkap dari login s.d. selesai'),
    ('5', 'revisi_output/diagrams/ERD.drawio',
     'Entity Relationship Diagram — 20 tabel dengan relasi 1:1 dan 1:N'),
    ('6', 'revisi_output/Bukti_Revisi_v2.docx',
     'Dokumen ini — bukti before/after seluruh poin revisi'),
], ['No', 'File', 'Keterangan'])

p('Catatan: Mockup PNG (364 file, ~total besar) sudah tersedia di bahan_skripsi/File-Figma-SPR-JTI/ '
  'dan direferensikan sebagai [Gambar 4.5]–[Gambar 4.9] di BAB IV. '
  'File PNG tidak disalin ke revisi_output karena ukurannya besar.',
  italic=True, size=10, space_after=4)

separator()
p('Dokumen ini dibuat otomatis pada 6 Oktober 2026.', size=10, italic=True,
  align=WD_ALIGN_PARAGRAPH.CENTER)

# ── Save ─────────────────────────────────────────────────────────────────────
doc.save(OUTPUT)
print(f'[OK] {OUTPUT}')
